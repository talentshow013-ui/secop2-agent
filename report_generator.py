import os
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

SHEET_MAP = {
    "suministro":    "Suministro",
    "eventos":       "Eventos y Logística",
    "social":        "Proyectos Sociales",
    "ambiental":     "Ambiental y Agrícola",
    "mantenimiento": "Mantenimiento",
    "otro":          "Otros",
}

COLUMNS = [
    ("ID Proceso",        "id_proceso",        22),
    ("Nombre del Proceso","nombre_proceso",     55),
    ("Entidad",           "entidad",            35),
    ("Valor (COP)",       "valor_proceso",      18),
    ("Fecha Publicación", "fecha_publicacion",  20),
    ("Departamento",      "departamento",       18),
    ("Ciudad",            "ciudad",             18),
    ("Score IA",          "_score",             10),
    ("Justificación IA",  "_justificacion",     45),
    ("URL SECOP 2",       "url_secop",          50),
    ("Fecha Scraped",     "_fecha_scraped",     20),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _write_header(ws):
    for col_idx, (header, _, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _get_or_create_wb(report_path: str) -> openpyxl.Workbook:
    if os.path.exists(report_path):
        return openpyxl.load_workbook(report_path)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEET_MAP.values():
        ws = wb.create_sheet(sheet_name)
        _write_header(ws)
    return wb


def append_processes(report_path: str, processes: list) -> None:
    if not processes:
        return

    wb = _get_or_create_wb(report_path)

    for proc in processes:
        categoria = proc.get("_categoria_claude") or proc.get("categoria", "otro")
        sheet_name = SHEET_MAP.get(categoria, SHEET_MAP["otro"])

        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            _write_header(ws)

        ws = wb[sheet_name]
        row_idx = ws.max_row + 1

        for col_idx, (_, field_key, _) in enumerate(COLUMNS, 1):
            if field_key == "_fecha_scraped":
                value = datetime.now().strftime("%Y-%m-%d %H:%M")
            else:
                value = proc.get(field_key, "")
                if value is None:
                    value = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx in [2, 9]))

            if field_key == "url_secop" and value:
                cell.hyperlink = str(value)
                cell.font = Font(color="0563C1", underline="single")

            if field_key == "_score":
                try:
                    score = int(value)
                    if score >= 80:
                        cell.fill = GREEN_FILL
                    elif score >= 60:
                        cell.fill = YELLOW_FILL
                except (ValueError, TypeError):
                    pass

    try:
        wb.save(report_path)
        log.info("Reporte Excel actualizado: %s (%d procesos)", report_path, len(processes))
    except Exception as e:
        log.error("Error guardando reporte Excel: %s", e)
